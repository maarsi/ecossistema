#!/home/mateus/env_transporte/bin/python
# -*- coding: utf-8 -*-
"""indice.py - indice semantico local de acervo documental (LanceDB + papel 'embeddings' via Ollama).

  indice.py indexar <pasta> --nome <colecao> [--incluir-codigo]
  indice.py buscar <colecao> "<pergunta>" [-k 6]
  indice.py validar <colecao> [-n 20]      valida com perguntas geradas sobre trechos sorteados
  indice.py listar

Manuais de software e referencias genericas: excluir com --excluir e indexar em colecao propria.
Tudo local. Cada trecho guarda arquivo e pagina. Reindexacao incremental por hash: so reprocessa o que mudou.
Formatos: PDF (PyMuPDF, por pagina), DOCX, DOC (via LibreOffice), MD, TXT; --incluir-codigo acrescenta .py e .sql.
Dados (CSV, XLSX, SHP, ZIP) ficam de fora de proposito. PDF sem camada de texto e listado como pendente de OCR.
"""
import argparse, hashlib, json, os, random, re, subprocess, sys, tempfile, time
from pathlib import Path
import requests

RAIZ = Path(__file__).resolve().parent.parent
BASE = Path("/mnt/nvme_lexar/datasets_sig/indices")
OLLAMA = os.environ.get("OLLAMA_HOST", "http://localhost:11434").rstrip("/")
MODELO = json.loads((RAIZ / "config" / "papeis.json").read_text(encoding="utf-8"))["embeddings"]["modelo"]
TAM, SOBRE, LOTE = 1200, 200, 32
EXCLUIR = ["_archive", "_to_delete", ".git", "__pycache__", ".venv"]
DOCS, CODIGO = {".pdf", ".docx", ".doc", ".md", ".txt"}, {".py", ".sql"}


def caminho_ubuntu(c):
    c = c.replace("\\", "/")
    return "/media/mateus/DADOS_01/" + c[3:] if c.lower().startswith("d:/") else c


def hash_arq(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()[:24]


def docx_texto(p):
    import docx
    d = docx.Document(p)
    partes = [x.text for x in d.paragraphs if x.text.strip()]
    for t in d.tables:
        for lin in t.rows:
            partes.append(" | ".join(c.text.strip() for c in lin.cells))
    return "\n".join(partes)


def extrair(p):
    """Lista de (pagina, texto). pagina = -1 quando o formato nao tem paginas."""
    s = p.suffix.lower()
    if s == ".pdf":
        import pymupdf
        with pymupdf.open(p) as d:
            return [(i + 1, pg.get_text()) for i, pg in enumerate(d)]
    if s == ".docx":
        return [(-1, docx_texto(p))]
    if s == ".doc":
        with tempfile.TemporaryDirectory() as tmp:
            subprocess.run(["soffice", "--headless", "--convert-to", "docx", "--outdir", tmp, str(p)],
                           capture_output=True, timeout=180)
            out = Path(tmp) / (p.stem + ".docx")
            return [(-1, docx_texto(out))] if out.exists() else []
    return [(-1, p.read_text(encoding="utf-8", errors="replace"))]


def fatiar(texto):
    texto = re.sub(r"[ \t]+", " ", texto).strip()
    if len(texto) <= TAM:
        return [texto] if len(texto) > 40 else []
    out, i = [], 0
    while i < len(texto):
        fim = min(len(texto), i + TAM)
        if fim < len(texto):
            q = texto.rfind("\n", i + TAM // 2, fim)
            fim = q if q > 0 else fim
        out.append(texto[i:fim].strip())
        if fim >= len(texto):
            break
        i = max(fim - SOBRE, i + 1)
    return [x for x in out if len(x) > 40]


def embed(textos):
    vs = []
    for k in range(0, len(textos), LOTE):
        r = requests.post(f"{OLLAMA}/api/embed", timeout=600,
                          json={"model": MODELO, "input": textos[k:k + LOTE], "keep_alive": "30m"}).json()
        if r.get("error"):
            sys.exit(f"ERRO embeddings: {r['error'][:150]}")
        vs += r["embeddings"]
    return vs


def abrir(nome):
    import lancedb
    return lancedb.connect(str(BASE / nome))


def tabela(registros):
    """Tabela Arrow com o vetor como lista de tamanho fixo (formato que o LanceDB reconhece para busca)."""
    import pyarrow as pa
    dim = len(registros[0]["vetor"])
    return pa.table({"arquivo": [r["arquivo"] for r in registros],
                     "pagina": pa.array([r["pagina"] for r in registros], pa.int32()),
                     "trecho": [r["trecho"] for r in registros],
                     "vetor": pa.array([r["vetor"] for r in registros], type=pa.list_(pa.float32(), dim))})


def q(s):
    return s.replace("'", "''")


def indexar(pasta, nome, incluir_codigo, excluir_extra=()):
    raiz = Path(caminho_ubuntu(pasta))
    if not raiz.exists():
        sys.exit(f"ERRO: pasta nao encontrada: {raiz}")
    exts = DOCS | (CODIGO if incluir_codigo else set())
    arqs = sorted(p for p in raiz.rglob("*") if p.is_file() and p.suffix.lower() in exts
                  and not any(x in p.relative_to(raiz).parts for x in EXCLUIR)
                  and not any(e.replace('\\', '/').strip('/') in p.relative_to(raiz).as_posix() for e in excluir_extra))
    BASE.mkdir(parents=True, exist_ok=True)
    db = abrir(nome)
    man = {}
    if "arquivos" in db.table_names():
        man = {r["arquivo"]: r for r in db.open_table("arquivos").to_arrow().to_pylist()}
    tbl = db.open_table("trechos") if "trechos" in db.table_names() else None
    t0, novos, pend, puladas, n_trechos = time.time(), [], [], 0, 0
    atuais = set()
    for p in arqs:
        rel = str(p.relative_to(raiz)); atuais.add(rel)
        h = hash_arq(p)
        if rel in man and man[rel]["hash"] == h:
            puladas += 1; continue
        try:
            paginas = extrair(p)
        except Exception as e:
            pend.append({"arquivo": rel, "hash": h, "n_trechos": 0, "status": f"erro {type(e).__name__}"}); continue
        total = sum(len(t.strip()) for _, t in paginas)
        if p.suffix.lower() == ".pdf" and total < 100 * max(1, len(paginas)):
            pend.append({"arquivo": rel, "hash": h, "n_trechos": 0, "status": "pendente_ocr"}); continue
        registros = [{"arquivo": rel, "pagina": pg, "trecho": tr} for pg, txt in paginas for tr in fatiar(txt)]
        if not registros:
            pend.append({"arquivo": rel, "hash": h, "n_trechos": 0, "status": "sem_texto"}); continue
        vetores = embed([f"Arquivo: {p.name}" + (f" | pagina {r['pagina']}" if r["pagina"] > 0 else "") + f"\n{r['trecho']}"
                         for r in registros])
        for r, v in zip(registros, vetores):
            r["vetor"] = v
        if tbl is not None:  # sempre substitui trechos antigos do arquivo: seguro contra interrupcao
            tbl.delete(f"arquivo = '{q(rel)}'")
        if tbl is None:
            tbl = db.create_table("trechos", data=tabela(registros))
        else:
            tbl.add(tabela(registros))
        novos.append({"arquivo": rel, "hash": h, "n_trechos": len(registros), "status": "ok"})
        n_trechos += len(registros)
        print(f"  + {rel[:90]} ({len(registros)} trechos)", flush=True)
    removidos = [a for a in man if a not in atuais]
    for a in removidos:
        if tbl is not None:
            tbl.delete(f"arquivo = '{q(a)}'")
    novo_man = {**{k: v for k, v in man.items() if k in atuais}, **{r["arquivo"]: r for r in novos + pend}}
    if novo_man:
        db.create_table("arquivos", data=list(novo_man.values()), mode="overwrite")
    (BASE / nome / "origem.txt").write_text(str(raiz), encoding="utf-8")
    print(f"INDICE '{nome}': {len(novos)} arquivos indexados ({n_trechos} trechos), {puladas} inalterados, "
          f"{len(removidos)} removidos, {len(pend)} pendentes, {time.time() - t0:.0f} s")
    for r in pend:
        print(f"  pendente: {r['arquivo'][:90]} [{r['status']}]")


def buscar(nome, pergunta, k=6, imprimir=True):
    db = abrir(nome)
    if "trechos" not in db.table_names():
        sys.exit(f"ERRO: indice '{nome}' vazio ou inexistente.")
    v = embed([pergunta])[0]
    busca = db.open_table("trechos").search(v, vector_column_name="vetor")
    try:
        busca = busca.distance_type("cosine")
    except AttributeError:
        busca = busca.metric("cosine")
    res = busca.limit(k).to_list()
    if imprimir:
        for i, r in enumerate(res, 1):
            pg = f", p. {r['pagina']}" if r["pagina"] > 0 else ""
            print(f"[{i}] {1 - r['_distance']:.3f} | {r['arquivo']}{pg}\n    {r['trecho'][:350].replace(chr(10), ' ')}\n")
    return res


def validar(nome, n):
    sys.path.insert(0, str(RAIZ / "bin")); import llm
    cfg = llm.PAPEIS["doc"]
    linhas = [r for r in abrir(nome).open_table("trechos").to_arrow().select(["arquivo", "pagina", "trecho"]).to_pylist()
              if len(r["trecho"]) > 500]
    random.seed(42); amostra = random.sample(linhas, min(n, len(linhas)))
    top1 = top5 = 0
    for r in amostra:
        pedido = ("Escreva UMA pergunta em portugues que o trecho abaixo responde. Use outras palavras, sem copiar frases, "
                  "sem numeros exatos e sem citar nomes de arquivo. Responda apenas com a pergunta.\n\nTRECHO:\n" + r["trecho"])
        perg, _, _ = llm.chamar_local(cfg["modelo"], pedido, 8192, 80, 0.2)
        perg = llm.limpar_raciocinio(perg).strip().splitlines()[0]
        res = buscar(nome, perg, 5, imprimir=False)
        pos = next((i for i, x in enumerate(res, 1) if x["arquivo"] == r["arquivo"] and x["pagina"] == r["pagina"]), None)
        top1 += pos == 1; top5 += pos is not None
        print(f"  {'OK ' if pos else 'ERR'} pos={pos} | {perg[:95]}")
    print(f"VALIDACAO '{nome}': origem em 1o lugar {top1}/{len(amostra)} | entre os 5 primeiros {top5}/{len(amostra)}")


def listar():
    if not BASE.exists():
        print("nenhum indice"); return
    for d in sorted(BASE.iterdir()):
        try:
            db = abrir(d.name); n = db.open_table("trechos").count_rows()
            origem = (d / "origem.txt").read_text() if (d / "origem.txt").exists() else "?"
            print(f"{d.name:24} {n:7} trechos | {origem}")
        except Exception:
            print(f"{d.name:24} (vazio)")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Indice semantico local do ecossistema")
    sub = ap.add_subparsers(dest="cmd", required=True)
    a1 = sub.add_parser("indexar"); a1.add_argument("pasta"); a1.add_argument("--nome", required=True)
    a1.add_argument("--incluir-codigo", action="store_true")
    a1.add_argument("--excluir", action="append", default=[], help="trecho de caminho a ignorar (repetivel)")
    a2 = sub.add_parser("buscar"); a2.add_argument("nome"); a2.add_argument("pergunta"); a2.add_argument("-k", type=int, default=6)
    a3 = sub.add_parser("validar"); a3.add_argument("nome"); a3.add_argument("-n", type=int, default=20)
    sub.add_parser("listar")
    a = ap.parse_args()
    if a.cmd == "indexar": indexar(a.pasta, a.nome, a.incluir_codigo, a.excluir)
    elif a.cmd == "buscar": buscar(a.nome, a.pergunta, a.k)
    elif a.cmd == "validar": validar(a.nome, a.n)
    else: listar()
